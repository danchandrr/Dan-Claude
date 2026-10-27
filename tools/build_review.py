import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from PIL import Image

d=json.load(open('review/data.json'))
D,MASK,READS,READ_AT,DDX,TIERS=d['D'],d['MASK'],d['READS'],d['READ_AT'],d['DDX'],d['TIERS']
PROCS,DISPO=d['PROCS'],d['DISPO_NAME']
NEW=set(d.get('NEW',[]))
POOLN={'paeds':'Paediatrics','obgyn':'Obstetrics & Gynaecology','cardio':'Cardiology','neuro':'Neurology','tox':'Toxicology','id':'Infectious Diseases','surg':'Surgery'}
KIND={'ecg':'12-lead ECG','abg':'Blood gas','xr':'X-ray','mic':'Microscopy'}
CAT={'h':'History','e':'Exam','l':'Tests','i':'Imaging'}
RHY={'sinus':'Sinus','afib':'AF','stemi':'ST elevation','peaked':'Peaked T','twinv':'T inversion','brugada':'Brugada','osborn':'Osborn waves','chb':'Complete heart block','avblock':'Dropped beats','alternans':'Electrical alternans','pericarditis':'Pericarditis','svt':'SVT','vt':'Ventricular tachycardia','torsades':'Torsades de pointes'}

TEAL=PatternFill('solid',fgColor='1C6973');NEWF=PatternFill('solid',fgColor='CDEFD8');YEL=PatternFill('solid',fgColor='FFF4C2');PALE=PatternFill('solid',fgColor='EEF4F3')
HEAD=Font(bold=True,color='FFFFFF');WRAP=Alignment(wrap_text=True,vertical='top');thin=Side(style='thin',color='C6D4D3')
BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
num=lambda a:'\n'.join(f'{i+1}. {x}' for i,x in enumerate(a)) if a else '—'

wb=Workbook()
# ---- How to review
ws=wb.active;ws.title='How to review'
rows=[('Zebra Shift: medical review sheet',),
('',),
('What is here',),
(f'• Cases: all {len(D)} patients, with every clue, the correct and wrong treatments, the destination, the teaching point, and how trick cases first present.',),
(f'• Pictures: all {len(READS)} ECGs, blood gases, X-rays and slides as the player sees them, with the answer options and the explanation shown afterwards.',),
('',),
('How to mark it',),
('1. Use the Status column: OK, Needs change, or leave it blank if you have not looked yet.',),
('2. Write what should change in the yellow "Your correction" column, in plain words. For example: "Tests clue 2: troponin should be 2,450 ng/L" or "Wrong options: remove IV heparin, it is defensible".',),
('3. You do not need to edit the clue cells themselves. Corrections are easier to apply from one column.',),
('4. Upload this file back to Claude when you are done (or part way). Every correction gets applied and checked in the game.',),
('',),
('Good to know',),
('• Cases are deliberately condensed for play: each clue is one line, and treatments are summarised.',),
('• The wrong treatment options should be clearly wrong for this patient, not just second-line.',),
('• "Trick presentation" is how some patients first walk in, looking like a simpler problem, before the real picture appears.',),
('• "Emergency differentials" are the other diagnoses offered when the case appears as a timed emergency.',),
('• Use the filter arrows in the header row to show one rarity, system or status at a time.',),
(f'• The {len(NEW)} cases with a green name were added most recently. Start with those if you have already reviewed the rest.',),
('• "Standards cross-match" lists every emergency in the five Indonesian specialist standards you sent and where it lives in the game. Green rows were added from those documents; red rows are not in the game yet, with the reason.',),
('• "Specialty pool" shows which specialist sees the case. Generalist specialties (EM, GP, Internal Medicine, Anaesthesia, Pathology) see every pool. "False alarm" cases arrive as emergencies but are benign.',),
]
for r in rows: ws.append(r)
ws.column_dimensions['A'].width=120
ws['A1'].font=Font(bold=True,size=16,color='1C6973')
for c in ('A3','A7','A13'): ws[c].font=Font(bold=True,size=12)
for row in ws.iter_rows(min_row=2): 
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')

status=DataValidation(type='list',formula1='"OK,Needs change"',allow_blank=True)

def sheet(name,headers,widths):
    w=wb.create_sheet(name);w.append(headers)
    for i,h in enumerate(headers,1):
        c=w.cell(row=1,column=i);c.fill=TEAL;c.font=HEAD;c.alignment=Alignment(wrap_text=True,vertical='center');c.border=BOX
        w.column_dimensions[c.column_letter].width=widths[i-1]
    w.row_dimensions[1].height=32;w.freeze_panes='C2'
    return w

# ---- Cases
H=['#','Disease','Rarity','System','Specialty pool','Age range','Sex','Presenting complaint','HR','BP','RR','SpO₂ %','Temp °C','Monitor rhythm',
   'History','Exam','Tests','Imaging','Shown as a picture first','Correct treatment','Wrong treatment options','Destination','Bedside procedures',
   'Also accepted as','Teaching point','Trick presentation','Looks like (mimic)','Emergency differentials','Status','Your correction']
W=[5,30,11,15,16,9,8,40,6,9,6,8,8,14,48,42,42,40,30,36,40,16,18,22,48,40,22,40,14,50]
ws=sheet('Cases',H,W);ws.add_data_validation(status)
order=sorted(range(len(D)),key=lambda i:(D[i]['t'],D[i]['n']))
for n,i in enumerate(order,1):
    x=D[i];v=x['v'];m=MASK.get(x['n'])
    pics='\n'.join(f"{CAT[c]} clue {k+1}: {KIND[READS[key]['k']]} ({READS[key]['o'][0]})" for c,k,key in READ_AT.get(x['n'],[])) or '—'
    trick=f"Walks in with: \"{m['cc']}\"\nVitals: HR {m['v'][0]}, BP {m['v'][1]}, RR {m['v'][2]}, SpO₂ {m['v'][3]}, T {m['v'][4]}\nThen: {'cardiac arrest' if m.get('turn')=='arrest' else 'crashes' if m.get('turn')=='crash' else 'stays the same'}" if m else '—'
    sex={'A':'Any','M':'Male','F':'Female'}[x['x']]
    row=[n,x['n'],TIERS[x['t']],x['s'],POOLN.get(x.get('sp'),'Core (everyone)')+(' · false alarm' if x.get('fake') else ''),f"{x['a'][0]}–{x['a'][1]}",sex,x['cc'],v[0],v[1],v[2],v[3],v[4],RHY.get(x.get('ecg') or 'sinus','Sinus'),
         num(x['h']),num(x['e']),num(x['l']),num(x['i']),pics,x['tx'],'\n'.join('• '+w for w in x['w']),DISPO.get(x.get('dispo') or '', x.get('dispo') or '—'),
         ', '.join(PROCS.get(p,p) for p in (x.get('proc') or [])) or '—',', '.join(x.get('al') or []) or '—',x['p'],trick,x.get('mim') or '—',', '.join(DDX.get(x['n'],[])) or '(picked automatically)','','']
    ws.append(row);r=ws.max_row
    for c in ws[r]: c.alignment=WRAP;c.border=BOX
    ws.cell(row=r,column=len(H)).fill=YEL
    if n%2==0:
        for c in ws[r][:len(H)-1]: c.fill=PALE
    if x['n'] in NEW: ws.cell(row=r,column=2).fill=NEWF;ws.cell(row=r,column=2).font=Font(bold=True)
    status.add(ws.cell(row=r,column=len(H)-1))
ws.auto_filter.ref=f"A1:{ws.cell(row=1,column=len(H)).column_letter}{ws.max_row}"
red=PatternFill('solid',fgColor='F9D6D3');grn=PatternFill('solid',fgColor='D5EDDC')
L=ws.cell(row=1,column=len(H)-1).column_letter
ws.conditional_formatting.add(f"{L}2:{L}{ws.max_row}",FormulaRule(formula=[f'{L}2="Needs change"'],fill=red))
ws.conditional_formatting.add(f"{L}2:{L}{ws.max_row}",FormulaRule(formula=[f'{L}2="OK"'],fill=grn))

# ---- Pictures
H2=['#','Picture ID','Type','Title shown','Used in case','Replaces this written clue','What the player sees','Correct answer','Other options shown','Explanation shown after answering','Status','Your correction']
W2=[5,12,12,24,28,40,84,38,40,50,14,50]
wp=sheet('Pictures',H2,W2);wp.freeze_panes='C2';wp.add_data_validation(status)
used={}
for dn,lst in READ_AT.items():
    x=next((y for y in D if y['n']==dn),None)
    for c,k,key in lst: used.setdefault(key,[]).append((dn,f"{CAT[c]} clue {k+1}: {x[c][k] if x else ''}"))
for n,key in enumerate(READS,1):
    r=READS[key];u=used.get(key,[])
    title=r.get('x2') or KIND[r['k']]
    wp.append([n,key,KIND[r['k']],title,'\n'.join(a for a,_ in u) or '—','\n'.join(b for _,b in u) or '—','',r['o'][0],'\n'.join('• '+o for o in r['o'][1:]),r['x'],'',''])
    row=wp.max_row
    for c in wp[row]: c.alignment=WRAP;c.border=BOX
    wp.cell(row=row,column=len(H2)).fill=YEL
    status.add(wp.cell(row=row,column=len(H2)-1))
    im=Image.open(f'review/pics/{key}.png');target=600 if r['k']=='ecg' else 300
    sc=target/im.width;xi=XLImage(f'review/pics/{key}.png');xi.width=int(im.width*sc);xi.height=int(im.height*sc)
    wp.add_image(xi,f"G{row}");wp.row_dimensions[row].height=max(60,xi.height*0.75+8)
L=wp.cell(row=1,column=len(H2)-1).column_letter
wp.conditional_formatting.add(f"{L}2:{L}{wp.max_row}",FormulaRule(formula=[f'{L}2="Needs change"'],fill=red))
wp.conditional_formatting.add(f"{L}2:{L}{wp.max_row}",FormulaRule(formula=[f'{L}2="OK"'],fill=grn))
wp.auto_filter.ref=f"A1:{wp.cell(row=1,column=len(H2)).column_letter}{wp.max_row}"
# ---- Standards cross-match: emergencies from the Indonesian specialist standards against the game's cases
import sys,os;sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from crossmatch import ROWS
TODAY=set(d.get('TODAY',[]))
BYN={x['n']:x for x in D}
miss=[c for r in ROWS for c in r[3] if c not in BYN]
assert not miss,('cross-match names not in the game',miss)
def where(n):
    x=BYN[n]
    if n in TODAY: return 'Added from these documents'
    return (POOLN.get(x.get('sp'))+' pool') if x.get('sp') else 'Core game'
HX=['Document','Section','Condition as written','Game case(s)','Where it is in the game','Note','Status','Your correction']
wx=sheet('Standards cross-match',HX,[30,30,40,44,34,40,14,40]);wx.add_data_validation(status)
GRN=PatternFill('solid',fgColor='CDEFD8');RED=PatternFill('solid',fgColor='F6D6D3')
for doc,sec,cond,cases,note in ROWS:
    wx.append([doc,sec,cond,'\n'.join(cases) or '—','\n'.join(sorted(set(where(c) for c in cases))) or 'Not in the game',note,'',''])
    r=wx.max_row
    for c in wx[r]: c.alignment=WRAP;c.border=BOX
    wx.cell(row=r,column=len(HX)).fill=YEL
    if not cases: wx.cell(row=r,column=5).fill=RED
    elif any(c in TODAY for c in cases): wx.cell(row=r,column=5).fill=GRN
    status.add(wx.cell(row=r,column=len(HX)-1))
wx.auto_filter.ref=f"A1:H{wx.max_row}"
import os;out=os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','review','zebra-shift-medical-review.xlsx')
import os;os.makedirs(os.path.dirname(out),exist_ok=True);wb.save(out);print('saved',out,os.path.getsize(out))
